"""Data of the three documents on the state of the repository at commit 102f399 -- 7plus1_status_2026-10-28 and its symbols and provenance
files -- which contain everything that the spreadsheet 7plus1_status_2026-10-28.xlsx contains.

Every number of the documents is read from tools/status_2026-10-28.json, the data file from which tools/make_status_xlsx.py writes the
spreadsheet.  The formulas of the spreadsheet are those of the workbook that make_status_xlsx.build makes from the same data; they are
evaluated here by a small evaluator of their own (SUM, COUNTA, COUNTIF, + and &), and the values are compared with the values that
LibreOffice computed and stored in the spreadsheet.  The checks of this module recompute from the repository itself whatever can be
recomputed: the statistics of the commits (git), the page counts and the changed documents (the PDF and Markdown files of the commits), the
presence of every quoted formula in its document at the commit described, and the outputs of the twelve check programs (the runs of
tools/make_docs.py)."""
import os, re, json, subprocess, hashlib, tempfile, datetime
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
NAME = '7plus1_status_2026-10-28'
DATA = os.path.join(HERE, 'status_2026-10-28.json')
XLSX = os.path.join(ROOT, NAME + '.xlsx')
SHEETS = ['Summary', 'Commits', 'Documents', 'Fixes', 'Checks', 'Private files']


def load():
    return json.load(open(DATA, encoding='utf-8'))


def present():
    return os.path.exists(DATA)


def sha256(path):
    h = hashlib.sha256()
    with open(path, 'rb') as f:
        for chunk in iter(lambda: f.read(1 << 20), b''): h.update(chunk)
    return h.hexdigest()


def git(*args, binary=False):
    """output of git in the repository (None if git or the repository is not available)"""
    try:
        p = subprocess.run(['git'] + list(args), cwd=ROOT, capture_output=True, timeout=120)
    except Exception:
        return None
    if p.returncode != 0:
        return None
    return p.stdout if binary else p.stdout.decode('utf-8', errors='replace')


def num(n):
    """an integer with thousands separators, as the spreadsheet shows it"""
    return '{:,}'.format(n) if isinstance(n, int) else str(n)


def when(t):
    """'2026-10-28 09:51' -> '28 October 2026, 09:51'"""
    d = datetime.datetime.strptime(t[:16], '%Y-%m-%d %H:%M')
    return '%d %s %d, %s' % (d.day, d.strftime('%B'), d.year, d.strftime('%H:%M'))


# the plain text of the data file -> the inline markup of the documents (mathematics, code, document names)
MATH = [('Θ^K, σ_r and Θ_K', '$\\Theta^K$, $\\sigma_r$ and $\\Theta_K$'), ('g_n = |dM_E/dφ|', '$g_n=|dM_E/d\\phi|$'),
        ('U^T h U', '$U^{\\mathsf T}hU$'), ('the Θ_K', 'the $\\Theta_K$'), ('matrix U of', 'matrix $U$ of'),
        ('in n dimensions', 'in $n$ dimensions'),
        ('CD = LC[g], RiemannCD, RicciCD, RicciScalarCD, EinsteinCD', '`CD = LC[g]`, `RiemannCD`, `RicciCD`, `RicciScalarCD`, `EinsteinCD`'),
        ('Code may break after / and _ and \\,', 'Code may break after `/`, `_` and `\\`,')]
FILE = re.compile(r'(?<![`\w])((?:tools/|build/|\.git/)[\w./*-]*\w|\.gitignore|tools/eval_log_7p1\.txt)')


def markup(text, S=None):
    """inline markup for a text of the data file: the formulas of MATH, file names and document names as code (code may break after _ and /)"""
    for a, b in MATH:
        text = text.replace(a, b)
    parts = re.split(r'(\$[^$]+\$|`[^`]+`)', text)
    names = sorted((d['name'] for d in (S or load())['documents']), key=len, reverse=True)
    pat = re.compile(r'(?<![\w`])(' + '|'.join(re.escape(n) for n in names) + r')(?![\w`])')
    for i, p in enumerate(parts):
        if p.startswith('$') or p.startswith('`'):
            continue
        p = FILE.sub(r'`\1`', p)
        sub = re.split(r'(`[^`]+`)', p)
        parts[i] = ''.join(s if s.startswith('`') else pat.sub(r'`\1`', s) for s in sub)
    return ''.join(parts)


# ------------------------------------------------------------------------------------------------------------------------------------------
# The formulas of the rewritten paragraphs, exactly as the documents write them.  'fix' is the row of the sheet Fixes (1-based), 'doc' the
# document, 'blocks' the displayed formulas (('eq', latex), ('align', lines) or ('code', text)), 'find' the strings that must occur in the
# Markdown file of the document at the commit described, 'words' an explanation of the symbols (every symbol is defined in the symbol file).
FORMULAS = [
    dict(fix=2, doc='README', title='The Clifford products $\\Theta^K$, the sign $\\sigma_r$ and the lowered $\\Theta_K$',
         lead='For the 256 increasing multi-indices $K=(A_1<\\dots<A_r)$, $r=0,\\dots,8$, of the eight real 16×16 Dirac matrices $\\Gamma^A$:',
         blocks=[('eq', r'\Theta^K=i^{\sigma_r}\Gamma^{A_1}\cdots\Gamma^{A_r},\qquad \sigma_r=\frac{r(r-1)}2\bmod2,\qquad \Theta_K=\eta_{A_1A_1}\cdots\eta_{A_rA_r}\Theta^K .')],
         find=[r'\Theta^K=i^{\sigma_r}\Gamma^{A_1}\cdots\Gamma^{A_r},\qquad \sigma_r=\frac{r(r-1)}2\bmod2,\qquad \Theta_K=\eta_{A_1A_1}\cdots\eta_{A_rA_r}\Theta^K'],
         words='The factor $i^{\\sigma_r}$ makes each $\\bar\\Psi\\Theta^K\\Psi$ real; $\\eta_{AA}$ are the diagonal entries of $\\eta=\\mathrm{diag}(+1,-1,\\dots,-1)$; '
               'for $r=0$ the product is $1_{16}$.'),
    dict(fix=3, doc='DIRAC16COMPLEX_PROVENANCE', title='The product and the left derivative of the Grassmann algebra',
         lead='For increasing index sets $K$, $K\'$ of the anticommuting generators $\\theta_k$ ($\\theta_k\\theta_l=-\\theta_l\\theta_k$), with $\\theta_K$ the '
              'product of the generators of $K$ in increasing order:',
         blocks=[('align', [r"\theta_K\theta_{K'}&=\mathrm{sign}(\pi)\,\theta_{\mathrm{sort}(K\cup K')}\qquad(\text{zero if }K\cap K'\neq\emptyset),",
                            r'\frac{\partial^L(\theta_{k_1}\cdots\theta_{k_r})}{\partial\theta_{k_j}}&=(-1)^{j-1}\,\theta_{k_1}\cdots\widehat{\theta_{k_j}}\cdots\theta_{k_r} .'])],
         find=[r"\theta_K\theta_{K'}&=\mathrm{sign}(\pi)\,\theta_{\mathrm{sort}(K\cup K')}\qquad(\text{zero if }K\cap K'\neq\emptyset)",
               r'\frac{\partial^L(\theta_{k_1}\cdots\theta_{k_r})}{\partial\theta_{k_j}}&=(-1)^{j-1}\,\theta_{k_1}\cdots\widehat{\theta_{k_j}}\cdots\theta_{k_r}'],
         words='$\\pi$ is the permutation that sorts the sequence $K$ followed by $K\'$, $\\mathrm{sort}(K\\cup K\')$ the sorted union, $\\partial^L$ the '
               'derivative from the left, and the hat marks the factor that is left out.'),
    dict(fix=4, doc='GRAVITY_PROVENANCE', title='The xAct definitions of the connection and the curvature tensors',
         lead='As the author’s notebook Einstein-Lovelock-7+1-Nash.nb writes them, after a component metric $g$ has been assigned (Wolfram Language with '
              'the packages xAct, xCoba and xTras):',
         blocks=[('code', 'CD = LC[g]; RiemannCD = Riemann[CD]; RicciCD = Ricci[CD];\nRicciScalarCD = RicciScalar[CD]; EinsteinCD = Einstein[CD];')],
         find=['CD = LC[g]; RiemannCD = Riemann[CD]; RicciCD = Ricci[CD];', 'RicciScalarCD = RicciScalar[CD]; EinsteinCD = Einstein[CD];'],
         words='After `CD = LC[g]` the abstract covariant derivative `CD` is the Levi-Civita connection of the component metric $g$, and '
               '`RiemannCD`, `RicciCD`, `RicciScalarCD`, `EinsteinCD` are its Riemann tensor, Ricci tensor, scalar curvature and Einstein tensor.'),
    dict(fix=5, doc='STANDARD_MODEL_PROVENANCE', title='The mixing angle of two neutrino flavours in matter (MSW)',
         lead='Diagonalizing the two-flavour Hamiltonian in matter gives the mixing angle in matter $\\theta_m$:',
         blocks=[('eq', r'\frac{\Delta m^2}{4E}\begin{pmatrix}-\cos2\theta&\sin2\theta\\\sin2\theta&\cos2\theta\end{pmatrix}+\mathrm{diag}(V,0)\quad\Longrightarrow\quad \tan2\theta_m=\frac{\sin2\theta}{\cos2\theta-2EV/\Delta m^2} .')],
         find=[r'\frac{\Delta m^2}{4E}\begin{pmatrix}-\cos2\theta&\sin2\theta\\\sin2\theta&\cos2\theta\end{pmatrix}+\mathrm{diag}(V,0)\quad\Longrightarrow\quad \tan2\theta_m=\frac{\sin2\theta}{\cos2\theta-2EV/\Delta m^2}'],
         words='$\\Delta m^2$ is the difference of the squared masses, $E$ the neutrino energy, $\\theta$ the mixing angle in vacuum and $V$ the potential '
               'of the electron neutrino in matter.'),
    dict(fix=6, doc='COUPLED_SOLUTIONS_PROVENANCE', title='The mode equation of pair creation',
         lead='The two-component amplitude $c(\\tau)$ of a mode of dirac16complex in the Einstein–Kasner solution of the earliest regime obeys',
         blocks=[('eq', r'i\,\frac{dc}{d\tau}=\omega\,(\cos\theta\,\sigma_1+\sin\theta\,\sigma_2)\,c,\qquad \omega=K\sqrt{\tau^{-2p_1}+\tau^{-2q_1}},\qquad \tan\theta=\tau^{p_1-q_1} .')],
         find=[r'i\,\frac{dc}{d\tau}=\omega\,(\cos\theta\,\sigma_1+\sin\theta\,\sigma_2)\,c,\qquad \omega=K\sqrt{\tau^{-2p_1}+\tau^{-2q_1}},\qquad \tan\theta=\tau^{p_1-q_1}'],
         words='$\\tau$ is the time since the singularity, $K$ the scaling variable of the mode, $p_1=(1+2\\sqrt2)/7$ and $q_1=u_1p_1$ with '
               '$u_1=-1+1/\\sqrt2$ the Kasner exponents of ordinary and auxiliary space, and $\\sigma_1$, $\\sigma_2$ Pauli matrices.'),
    dict(fix=7, doc='COUPLED_SOLUTIONS', title='The minisuperspace action and the identities of its variation',
         lead='For the metric $ds^2=n^2dx_0^2-a^2(dx_1^2+dx_2^2+dx_3^2)-b^2(dx_4^2+\\dots+dx_7^2)$, with $\\mathrm{EL}_q$ the Euler–Lagrange expression for '
              'the variable $q$, evaluated at $n=1$:',
         blocks=[('align', [r'S&=-\frac1{2\kappa}\int n\,a^3b^4\sum_k\alpha_kL_k\,dx_0-\int nW\,dx_0,',
                            r'\mathrm{EL}_n(na^3b^4L_k)&=-2a^3b^4E^{(k)0}{}_0,\qquad \mathrm{EL}_a(na^3b^4L_k)=-6a^2b^4E^{(k)1}{}_1,',
                            r'\mathrm{EL}_b(na^3b^4L_k)&=-8a^3b^3E^{(k)4}{}_4 .'])],
         find=[r'S&=-\frac1{2\kappa}\int n\,a^3b^4\sum_k\alpha_kL_k\,dx_0-\int nW\,dx_0',
               r'\mathrm{EL}_n(na^3b^4L_k)&=-2a^3b^4E^{(k)0}{}_0,\qquad \mathrm{EL}_a(na^3b^4L_k)=-6a^2b^4E^{(k)1}{}_1',
               r'\mathrm{EL}_b(na^3b^4L_k)&=-8a^3b^3E^{(k)4}{}_4'],
         words='$\\kappa=8\\pi G_8$, $\\alpha_k$ are the Lovelock couplings ($\\alpha_0=2\\Lambda$, $\\alpha_1=1$), $L_k$ the Lovelock densities, $E^{(k)A}{}_B$ the '
               'Lovelock tensors and $W(a,b)$ the comoving energy of the matter.'),
    dict(fix=8, doc='COUPLED_SOLUTIONS_PROVENANCE', title='The comoving energy of one massless particle; the minisuperspace action and its identities',
         lead='For one massless particle with comoving momenta $k_a$ along ordinary space and $k_b$ along the auxiliary space:',
         blocks=[('eq', r'W=\omega=\sqrt{k_a^2/a^2+k_b^2/b^2},\qquad \partial_aW=-\frac{k_a^2}{a^3\omega},\qquad \partial_bW=-\frac{k_b^2}{b^3\omega} .'),
                 ('align', [r'S&=-\frac1{2\kappa}\int na^3b^4\sum_k\alpha_kL_k\,dx_0-\int nW\,dx_0,',
                            r'\mathrm{EL}_n(na^3b^4L_k)&=-2a^3b^4E^{(k)0}{}_0,\qquad \mathrm{EL}_a(na^3b^4L_k)=-6a^2b^4E^{(k)1}{}_1,',
                            r'\mathrm{EL}_b(na^3b^4L_k)&=-8a^3b^3E^{(k)4}{}_4 .'])],
         find=[r'W=\omega=\sqrt{k_a^2/a^2+k_b^2/b^2},\qquad \partial_aW=-\frac{k_a^2}{a^3\omega},\qquad \partial_bW=-\frac{k_b^2}{b^3\omega}',
               r'S&=-\frac1{2\kappa}\int na^3b^4\sum_k\alpha_kL_k\,dx_0-\int nW\,dx_0'],
         words='The second display is the minisuperspace action with the identities of its variation, the same equations as those of the '
               'previous paragraph, except that the provenance file writes $na^3b^4$ without the thin space of $n\\,a^3b^4$; $\\omega$ is the energy '
               'of the particle, and the derivatives $\\partial_aW$, $\\partial_bW$ enter the variations with respect to $a$ and $b$.'),
    dict(fix=9, doc='DARK_SECTOR', title='The Einstein-frame mass of a Kaluza–Klein level and its coupling to the radion',
         lead='A level of Kaluza–Klein mass $\\mu_n$ of dirac16complex (mass $m$) has, in the Einstein frame, the mass $M_E(\\beta)$, with $M_n=M_E(0)$, '
              'and hence the coupling $g_n$ to the canonical radion $\\phi=\\sqrt{12}M_{\\rm Pl}\\beta$:',
         blocks=[('eq', r'M_E(\beta)=e^{-2\beta}\sqrt{m^2+\mu_n^2e^{-2\beta}},\qquad g_n=\Big|\frac{dM_E}{d\phi}\Big|_{\beta=0}=\frac{2M_n+\mu_n^2/M_n}{\sqrt{12}\,M_{\rm Pl}} .')],
         find=[r'M_E(\beta)=e^{-2\beta}\sqrt{m^2+\mu_n^2e^{-2\beta}},\qquad g_n=\Big|\frac{dM_E}{d\phi}\Big|_{\beta=0}=\frac{2M_n+\mu_n^2/M_n}{\sqrt{12}\,M_{\rm Pl}}'],
         words='$\\beta=\\ln(b/b_0)$ is the radion and $M_{\\rm Pl}$ the reduced Planck mass.'),
    dict(fix=10, doc='DIRAC16COMPLEX', title='The five momenta of the exact mode sum',
         lead='7+1.nb sums the mode algebra exactly, with symbolic $\\vec x,\\vec y$ and $m=1$, for the symmetric set of five momenta',
         blocks=[('eq', r'\vec p\in\big\{\pm(2,2,4,0,0,0,0),\ \pm(0,0,0,2,2,0,0),\ \vec 0\big\}\qquad(E=5,5,3,3,1).')],
         find=[r'\vec p\in\big\{\pm(2,2,4,0,0,0,0),\ \pm(0,0,0,2,2,0,0),\ \vec 0\big\}\qquad(E=5,5,3,3,1)'],
         words='$E=\\sqrt{m^2+|\\vec p|^2}$: $\\sqrt{1+24}=5$, $\\sqrt{1+8}=3$, $\\sqrt1=1$.'),
    dict(fix=11, doc='DIRAC16COMPLEX_PROVENANCE', title='The orthogonal matrix $U$ of the spinors',
         lead='With the real spinors $u_s(\\vec p)$ of positive and $v_s(-\\vec p)$ of negative energy ($s=1,\\dots,8$, $u_s^{\\mathsf T}u_s=2E$) and the real '
              'symmetric plane-wave Dirac Hamiltonian $h(\\vec p)$:',
         blocks=[('eq', r'U=\frac1{\sqrt{2E}}\big(u_1(\vec p),\dots,u_8(\vec p),v_1(-\vec p),\dots,v_8(-\vec p)\big)\qquad\text{satisfies}\qquad U^{\mathsf T}h(\vec p)\,U=\mathrm{diag}(E1_8,-E1_8).')],
         find=[r'U=\frac1{\sqrt{2E}}\big(u_1(\vec p),\dots,u_8(\vec p),v_1(-\vec p),\dots,v_8(-\vec p)\big)\qquad\text{satisfies}\qquad U^{\mathsf T}h(\vec p)\,U=\mathrm{diag}(E1_8,-E1_8)'],
         words='$U$ is real and orthogonal, $U^{\\mathsf T}U=1_{16}$; $h(\\vec p)=p^i\\alpha_i+m\\Gamma^0$ with $\\alpha_i=\\Gamma^0\\Gamma^i$ ($i=1,\\dots,7$, summed).'),
    dict(fix=14, doc='KOHN_SHAM_GRAVITY_PROVENANCE', title='The completeness of the $\\Theta_K$ and the Fierz identity',
         lead='The 256 matrices $\\Theta_K$ are complete (the first display), and this gives the Fierz identity (the second display) for commuting '
              'spinors $a,b,c,d$:',
         blocks=[('eq', r'\sum_K(\Theta_K)_{cd}(\Theta^K)_{ab}=16\,\delta_{ad}\delta_{bc},'),
                 ('eq', r'\sum_{K\in r}(a\Theta_Kb)(c\Theta^Kd)=\sum_JC_{rJ}\,(a\Theta_Jd)(c\Theta^Jb),\qquad C_{rJ}=\frac1{256}\sum_{K\in r}\mathrm{tr}\big(\Theta_K\Theta_J\Theta^K\Theta^J\big).')],
         find=[r'\sum_K(\Theta_K)_{cd}(\Theta^K)_{ab}=16\,\delta_{ad}\delta_{bc}',
               r'\sum_{K\in r}(a\Theta_Kb)(c\Theta^Kd)=\sum_JC_{rJ}\,(a\Theta_Jd)(c\Theta^Jb),\qquad C_{rJ}=\frac1{256}\sum_{K\in r}\mathrm{tr}\big(\Theta_K\Theta_J\Theta^K\Theta^J\big)'],
         words='Here $(aMb)=\\sum_{i,j}a_iM_{ij}b_j$ for a 16×16 matrix $M$; $K\\in r$ runs over the multi-indices of rank $r$ and $J$ over all 256; as '
               'subscripts in the first display, $a,b,c,d$ are spinor indices $1,\\dots,16$, and $\\delta$ is the Kronecker delta.'),
    dict(fix=15, doc='README_PROVENANCE', title='The notation of sections 7–10',
         lead='With $\\Psi$ the field dirac16complex:',
         blocks=[('eq', r'\Gamma_A=\Gamma^{(1,7)}_A,\qquad \Gamma^{A}=\eta_{(1,7)}^{AA}\Gamma_A,\qquad \Gamma^0=\Gamma_0,\qquad \chi=\chi^{(1,7)}=i\Gamma^{(1,7)}_8,\\ \bar\Psi=\Psi^\dagger\Gamma^0,\qquad \vec x=(x_1,\dots,x_7),\qquad \vec p=(p^1,\dots,p^7).')],
         find=[r'\Gamma_A=\Gamma^{(1,7)}_A,\qquad \Gamma^{A}=\eta_{(1,7)}^{AA}\Gamma_A,\qquad \Gamma^0=\Gamma_0,\qquad \chi=\chi^{(1,7)}=i\Gamma^{(1,7)}_8,\\ \bar\Psi=\Psi^\dagger\Gamma^0'],
         words='$\\Gamma^{(1,7)}_A$ are the eight real 16×16 Dirac matrices of 7+1.nb, $\\eta_{(1,7)}=\\mathrm{diag}(+1,-1,\\dots,-1)$, and '
               '$\\Gamma^{(1,7)}_8=\\Gamma^{(1,7)}_0\\cdots\\Gamma^{(1,7)}_7$.'),
    dict(fix=16, doc='STANDARD_MODEL', title='The electroweak scale and the masses at tree level',
         lead='The reworded sentence states:',
         blocks=[('eq', r"v=(\sqrt2G_F)^{-1/2}=246.2196402\ \text{GeV},\qquad m_W=\frac{gv}2,\qquad m_Z=\frac{\sqrt{g^2+g'^2}\,v}2,\qquad m_H^2=2\lambda v^2 .")],
         find=[r'$`v=(\sqrt2G_F)^{-1/2}=246.2196402`$ GeV',
               r"At tree level the masses are $`m_W=gv/2`$, $`m_Z=\sqrt{g^2+g'^2}\,v/2`$ and $`m_H^2=2\lambda v^2`$"],
         words='$G_F$ is the Fermi constant, $g$ and $g\'$ the gauge couplings of $SU(2)_L$ and $U(1)_Y$, $\\lambda$ the quartic coupling of the Higgs field; '
               'the document writes the three masses inline, as $m_W=gv/2$ and $m_Z=\\sqrt{g^2+g\'^2}\\,v/2$.'),
    dict(fix=17, doc='STANDARD_MODEL_SYMBOLS', title='The matrix notation',
         lead='The reworded convention:',
         blocks=[('eq', r'1_n=(\delta_{ij})_{i,j=1}^n,\qquad (M^\dagger)_{ij}=\overline{M_{ji}},\qquad (M^*)_{ij}=\overline{M_{ij}},\\ (A\otimes B)_{(i-1)p+k,\,(j-1)q+l}=A_{ij}B_{kl}\ \ (B\ \text{of size}\ p\times q).')],
         find=[r'$`1_n`$ is the $`n\times n`$ identity, $`M^\dagger`$ the conjugate transpose, $`M^*`$ the complex conjugate and $`A\otimes B`$ the '
               r'Kronecker product (in Mathematica `KroneckerProduct[A, B]`)'],
         words='$1_n$ is the $n\\times n$ identity, $M^\\dagger$ the conjugate transpose, $M^*$ the complex conjugate (the bar is complex conjugation of a '
               'number) and $A\\otimes B$ the Kronecker product, in Mathematica `KroneckerProduct[A, B]`; the document states these in words, the display '
               'writes them out.'),
]


def emit_formulas(d, heading=None, level=2):
    """write every formula of the rewritten paragraphs into the document d: a heading per paragraph (of the given level), the lead, the
    displays, the explanation"""
    for k, F in enumerate(FORMULAS):
        if heading:
            (d.h2 if level == 2 else d.h3)(heading(k, F))
        d.p(F['lead'])
        for b in F['blocks']:
            if b[0] == 'eq': d.eq(b[1])
            elif b[0] == 'align': d.align(b[1])
            else: d.code(b[1], 'mathematica')
        d.p(F['words'])


def typographic(d):
    """straight double quotes outside code spans and inline mathematics -> typographic quotes (an opening one after a space, a bracket or
    a dash, a closing one otherwise), in the title, the abstract and every block of text and table of the document d"""
    def fix(s):
        if not isinstance(s, str) or '"' not in s:
            return s
        out = ''
        for part in re.split(r'(`[^`]+`|\$[^$]+\$)', s):
            if part.startswith('`') or part.startswith('$'):
                out += part; continue
            for ch in part:
                if ch == '"':
                    prev = out[-1] if out else ' '
                    out += '“' if (prev.isspace() or prev in '([{—–/') else '”'
                else:
                    out += ch
        return out
    d.title = fix(d.title); d.subtitle = fix(d.subtitle); d.abstract = fix(d.abstract); d.meta = [fix(m) for m in d.meta]
    blocks = []
    for b in d.blocks:
        k = b[0]
        if k in ('h1', 'h2', 'h3', 'p', 'quote'):
            b = (k, fix(b[1]))
        elif k in ('ul', 'ol'):
            b = (k, [fix(x) for x in b[1]])
        elif k == 'table':
            b = (k, [fix(h) for h in b[1]], [[fix(c) for c in r] for r in b[2]], b[3], b[4])
        blocks.append(b)
    d.blocks = blocks
    return d


def formula_presence(commit):
    """for each formula: is every string of F['find'] in the Markdown file of its document at the commit (git show), or, without git, in the
    file on disk?"""
    out = []
    for F in FORMULAS:
        txt = git('show', '%s:%s.md' % (commit, F['doc']))
        src = 'git show %s:%s.md' % (commit, F['doc'])
        if txt is None:
            p = os.path.join(ROOT, F['doc'] + '.md')
            txt = open(p, encoding='utf-8').read() if os.path.exists(p) else ''
            src = F['doc'] + '.md (on disk)'
        found = [s in txt for s in F['find']]
        out.append((F, all(found), len(found), sum(found), src))
    return out


# ------------------------------------------------------------------------------------------------------------------------------------------
# A small evaluator of the formulas of the spreadsheet (the functions SUM, COUNTA, COUNTIF and the operators + and &, as the spreadsheet uses
# them), working on the workbook that make_status_xlsx.build makes in memory.
REF = re.compile(r"(?:(?:'([^']+)'|([A-Za-z][A-Za-z ]*))!)?\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?")


def _col(s):
    n = 0
    for ch in s: n = n * 26 + ord(ch) - 64
    return n


def _split_top(expr, ops):
    """split expr at the operators ops that are outside parentheses and strings; returns [(op, term), ...]"""
    parts = []; depth = 0; instr = False; cur = ''; op = None
    for ch in expr:
        if ch == '"': instr = not instr
        if not instr:
            if ch == '(': depth += 1
            elif ch == ')': depth -= 1
            elif depth == 0 and ch in ops:
                parts.append((op, cur)); cur = ''; op = ch; continue
        cur += ch
    parts.append((op, cur))
    return parts


class Evaluator:
    def __init__(self, wb):
        self.wb = wb

    def cells(self, sheet, ref):
        m = REF.fullmatch(ref.strip())
        if not m:
            raise ValueError('not a reference: ' + ref)
        sh = m.group(1) or m.group(2) or sheet
        c1, r1 = _col(m.group(3)), int(m.group(4))
        c2, r2 = (_col(m.group(5)), int(m.group(6))) if m.group(5) else (c1, r1)
        return sh, [(sh, r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]

    def value(self, sh, r, c):
        v = self.wb[sh].cell(row=r, column=c).value
        if isinstance(v, str) and v.startswith('='):
            return self.formula(sh, v)
        return v

    def formula(self, sheet, f):
        return self.expr(sheet, f[1:] if f.startswith('=') else f)

    def expr(self, sheet, e):
        parts = _split_top(e, '&+')
        val = None
        for op, term in parts:
            v = self.term(sheet, term.strip())
            if op is None: val = v
            elif op == '+': val = val + v
            elif op == '&': val = self.text(val) + self.text(v)
        return val

    @staticmethod
    def text(v):
        if isinstance(v, float) and v.is_integer(): v = int(v)
        return '' if v is None else str(v)

    def term(self, sheet, t):
        if t.startswith('"') and t.endswith('"'):
            return t[1:-1]
        if re.fullmatch(r'-?\d+(\.\d+)?', t):
            return float(t) if '.' in t else int(t)
        m = re.fullmatch(r'(SUM|COUNTA|COUNTIF)\((.*)\)', t)
        if m:
            fn, args = m.group(1), [a for _, a in _split_top(m.group(2), ',')]
            sh, refs = self.cells(sheet, args[0])
            vals = [self.value(*x) for x in refs]
            if fn == 'SUM':
                return sum(v for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool))
            if fn == 'COUNTA':
                return sum(1 for v in vals if v not in (None, ''))
            crit = self.term(sheet, args[1].strip())
            if isinstance(crit, str):
                return sum(1 for v in vals if isinstance(v, str) and v.lower() == crit.lower())
            return sum(1 for v in vals if isinstance(v, (int, float)) and not isinstance(v, bool) and v == crit)
        sh, refs = self.cells(sheet, t)
        if len(refs) != 1:
            raise ValueError('a range where one cell is expected: ' + t)
        return self.value(*refs[0])


def workbook(S):
    """the workbook that tools/make_status_xlsx.py writes, made in memory from the same data"""
    import importlib.util
    spec = importlib.util.spec_from_file_location('make_status_xlsx', os.path.join(HERE, 'make_status_xlsx.py'))
    mod = importlib.util.module_from_spec(spec); spec.loader.exec_module(mod)
    return mod.build(S)


def formulas(S):
    """every formula of the spreadsheet: (sheet, cell, formula, value computed here), in the order of the sheets"""
    wb = workbook(S); ev = Evaluator(wb)
    out = []
    for name in SHEETS:
        ws = wb[name]
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    out.append((name, c.coordinate, c.value, ev.formula(name, c.value)))
    return out, wb


def summary_rows(S):
    """the rows of the sheet Summary: (item, value computed here, where it comes from, formula)"""
    fs, wb = formulas(S)
    val = {(s, c): v for s, c, f, v in fs}
    ws = wb['Summary']; rows = []
    r = 5
    while ws.cell(row=r, column=1).value:
        f = ws.cell(row=r, column=2).value
        rows.append((ws.cell(row=r, column=1).value, val[('Summary', 'B%d' % r)], ws.cell(row=r, column=3).value, f))
        r += 1
    return rows


def stored_values():
    """the values that LibreOffice computed and stored in the spreadsheet, {(sheet, cell): value}; None if the file is missing"""
    if not os.path.exists(XLSX):
        return None
    from openpyxl import load_workbook
    wb = load_workbook(XLSX, data_only=True)
    return {(ws.title, c.coordinate): c.value for ws in wb.worksheets for row in ws.iter_rows() for c in row}


def same(a, b):
    if isinstance(a, (int, float)) and isinstance(b, (int, float)) and not isinstance(a, bool) and not isinstance(b, bool):
        return abs(a - b) < 1e-9
    return Evaluator.text(a) == Evaluator.text(b)


def stored_comparison(S):
    """(number of formulas, number whose stored value equals the value computed here, list of differences); None if no spreadsheet"""
    st = stored_values()
    if st is None:
        return None
    fs, wb = formulas(S)
    diff = [(s, c, f, v, st.get((s, c))) for s, c, f, v in fs if not same(v, st.get((s, c)))]
    return len(fs), len(fs) - len(diff), diff


def xlsx_application():
    """the application recorded in the spreadsheet (docProps/app.xml)"""
    if not os.path.exists(XLSX):
        return ''
    import zipfile
    try:
        app = zipfile.ZipFile(XLSX).read('docProps/app.xml').decode('utf-8', errors='replace')
    except Exception:
        return ''
    m = re.search(r'<Application>([^<]*)</Application>', app)
    return m.group(1).replace('$', ' ') if m else ''


# ------------------------------------------------------------------------------------------------------------------------------------------
# Checks of the data file against the repository
def commit_checks(S):
    """for each commit: (hash, full name agrees, (files, added, removed) agree, the values from git); None if git is not available"""
    out = []
    for c in S['commits']:
        full = git('rev-parse', c['hash'])
        st = git('show', '--shortstat', '--format=', c['hash'])
        if full is None or st is None:
            return None
        g = lambda pat: int(re.search(pat, st).group(1)) if re.search(pat, st) else 0
        vals = (g(r'(\d+) files? changed'), g(r'(\d+) insertions?\(\+\)'), g(r'(\d+) deletions?\(-\)'))
        out.append((c['hash'], full.strip() == c['full'], vals == (c['files'], c['added'], c['removed']), vals))
    return out


def page_checks(S):
    """for each document: (name, pages in the data file, pages of the PDF at the commit described); None without git or pdfinfo"""
    commit = S['described_commit']
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        for d in S['documents']:
            pdf = git('show', '%s:%s.pdf' % (commit, d['name']), binary=True)
            if pdf is None:
                return None
            p = os.path.join(tmp, 'x.pdf')
            open(p, 'wb').write(pdf)
            try:
                info = subprocess.run(['pdfinfo', p], capture_output=True, text=True, timeout=60).stdout
            except Exception:
                return None
            m = re.search(r'^Pages:\s+(\d+)', info, re.M)
            out.append((d['name'], d['pages'], int(m.group(1)) if m else None))
    return out


def changed_checks(S):
    """for the commits 93c500f and 102f399: (commit, documents marked changed, documents whose files the commit changes, agreement)"""
    out = []
    for c, key in (('93c500f', 'changed_93c500f'), ('102f399', 'changed_102f399')):
        names = git('show', '--name-only', '--format=', c)
        if names is None:
            return None
        names = set(names.split())
        actual = [d['name'] for d in S['documents'] if any(d['name'] + e in names for e in ('.md', '.tex', '.pdf'))]
        marked = [d['name'] for d in S['documents'] if d[key]]
        out.append((c, len(marked), len(actual), marked == actual))
    return out


TRUE_LINE = re.compile(r'(:\s*True\s*$)|(\(True\)\s*$)')
FALSE_LINE = re.compile(r'(:\s*False\s*$)|(\(False\)\s*$)')


def program_checks(S, runs):
    """for each program: (program, data (exit, True, False), recounted from the runs of this build, agreement); None without runs"""
    if not runs:
        return None
    by = {r[0].split(' ')[0]: r for r in runs}
    out = []
    for p in S['programs']:
        r = by.get(p['program'])
        if r is None:
            out.append((p['program'], (p['exit_code'], p['true'], p['false']), None, False)); continue
        lines = r[1].strip().splitlines()
        m = re.search(r'exit code (-?\d+)', r[0])          # the title of a run ends in "...; exit code N)"
        rc = r[2] if len(r) > 2 else (int(m.group(1)) if m else None)
        got = (rc, sum(1 for l in lines if TRUE_LINE.search(l)), sum(1 for l in lines if FALSE_LINE.search(l)))
        want = (p['exit_code'], p['true'], p['false'])
        out.append((p['program'], want, got, got == want))
    return out
