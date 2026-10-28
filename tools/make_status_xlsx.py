#!/usr/bin/env python3
"""Build the spreadsheet 7plus1_status_2026-10-28.xlsx from tools/status_2026-10-28.json.

    python3 tools/make_status_xlsx.py            -> writes ./7plus1_status_2026-10-28.xlsx

The spreadsheet records the state of the repository at commit 102f399 (28 October 2026): the commits of that day, the 30
documents with their pages and loose lines, the paragraphs that were rewritten, the outputs of the twelve check programs, the
checks of the notebook, the comparison of a fresh clone with the commit, and the private files that are kept out of the repository.
Every number is taken from the data file; every total and every value of the sheet Summary is a formula.  openpyxl writes the
formulas without computed values; when LibreOffice (soffice) is installed the file is then opened and saved again by it in headless
mode, which computes and stores every value (Excel and LibreOffice also compute them on opening).  Finally the computed values are
read back and checked: no formula may evaluate to an error.
"""
import json, os, shutil, subprocess, sys, tempfile, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(HERE)
DATA = os.path.join(HERE, 'status_2026-10-28.json')
OUT = os.path.join(ROOT, '7plus1_status_2026-10-28.xlsx')

F = 'Arial'
FONT = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
HFONT = Font(name=F, size=10, bold=True, color='FFFFFF')
TITLE = Font(name=F, size=14, bold=True, color='1F3B4D')
SUB = Font(name=F, size=10, italic=True, color='4A5A66')
NOTE = Font(name=F, size=9, italic=True, color='4A5A66')
GREY = Font(name=F, size=10, color='4A5A66')
HFILL = PatternFill('solid', fgColor='1F4E5F')
TOTAL = PatternFill('solid', fgColor='E4EEF1')
BLOCK = PatternFill('solid', fgColor='D5E5EA')
TOP = Border(top=Side(style='thin', color='1F4E5F'))
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='top')
LEFT = Alignment(horizontal='left', vertical='top', wrap_text=True)


def header(ws, row, names, widths):
    for j, n in enumerate(names, 1):
        c = ws.cell(row=row, column=j, value=n)
        c.font = HFONT; c.fill = HFILL; c.alignment = Alignment(wrap_text=True, vertical='center')
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w


def put(ws, row, values, fmt=None, bold=False):
    for j, v in enumerate(values, 1):
        c = ws.cell(row=row, column=j, value=v)
        c.font = BOLD if bold else FONT
        c.alignment = WRAP if isinstance(v, str) and not v.startswith('=') else Alignment(vertical='top')
        if fmt and j in fmt:
            c.number_format = fmt[j]


def total_style(ws, row, ncol, centered_from=None):
    for j in range(1, ncol + 1):
        c = ws.cell(row=row, column=j); c.fill = TOTAL; c.border = TOP
        if centered_from and j >= centered_from:
            c.alignment = CENTER


def note(ws, row, text, ncol):
    ws.cell(row=row, column=1, value=text).font = NOTE
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    ws.cell(row=row, column=1).alignment = Alignment(wrap_text=True, vertical='top')
    width = sum((ws.column_dimensions[get_column_letter(j)].width or 10) for j in range(1, ncol + 1))
    lines = max(1, -(-int(len(text) * 1.05) // int(width)))          # a merged cell does not grow by itself
    ws.row_dimensions[row].height = 13 * lines + 4


def block_title(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = BOLD
    ws.cell(row=row, column=1).fill = BLOCK; ws.cell(row=row, column=2).fill = BLOCK


def build(S):
    wb = Workbook()
    dc = S['described_commit']

    # ---------------------------------------------------------------- Commits
    wc = wb.active; wc.title = 'Commits'
    header(wc, 1, ['Commit', 'Date and time (PDT)', 'Title', 'What it contains', 'Files changed', 'Lines added', 'Lines removed'],
           [10, 17, 52, 52, 12, 12, 13])
    for i, c in enumerate(S['commits'], 2):
        put(wc, i, [c['hash'], datetime.datetime.strptime(c['time'], '%Y-%m-%d %H:%M'), c['title'], c['contents'], c['files'], c['added'],
                    c['removed']], fmt={2: 'yyyy-mm-dd hh:mm', 5: '#,##0', 6: '#,##0', 7: '#,##0'})
    cl = len(S['commits']) + 1
    put(wc, cl + 1, ['Total', None, None, None] + ['=SUM(%s2:%s%d)' % (L, L, cl) for L in 'EFG'], fmt={5: '#,##0', 6: '#,##0', 7: '#,##0'},
        bold=True)
    total_style(wc, cl + 1, 7)
    note(wc, cl + 3, 'Source: git log and git show --shortstat of %s, branch %s. A file changed in two commits is counted in both. Commit %s '
         'corrects two check counts that were found while this spreadsheet was being prepared; the spreadsheet and its documents were added in '
         'the next commit.' % (S['repository'], S['branch'], dc), 7)
    wc.freeze_panes = 'A2'

    # ---------------------------------------------------------------- Documents
    wd = wb.create_sheet('Documents')
    header(wd, 1, ['Document', 'Kind', 'Covers (sections of 7+1.nb)', 'Pages', 'Loose lines, first build of 28 Oct', 'Loose lines after 0654dcd',
                   'Loose lines now', 'Overflowing lines now', 'LaTeX errors now', 'Changed by 93c500f', 'Changed by 102f399'],
           [38, 11, 24, 8, 14, 13, 10, 12, 10, 11, 11])
    for i, d in enumerate(S['documents'], 2):
        put(wd, i, [d['name'], d['kind'], d['covers'], d['pages'], d['loose_start'], d['loose_after_0654dcd'], d['loose_now'], d['overfull_now'],
                    d['errors_now'], 'Yes' if d['changed_93c500f'] else 'No', 'Yes' if d['changed_102f399'] else 'No'])
        for j in range(4, 12):
            wd.cell(row=i, column=j).alignment = CENTER
    dl = len(S['documents']) + 1
    put(wd, dl + 1, ['Total (%d documents)' % len(S['documents']), None, None] + ['=SUM(%s2:%s%d)' % (L, L, dl) for L in 'DEFGHI'] +
        ['=COUNTIF(J2:J%d,"Yes")' % dl, '=COUNTIF(K2:K%d,"Yes")' % dl], fmt={4: '#,##0'}, bold=True)
    total_style(wd, dl + 1, 11, centered_from=4)
    note(wd, dl + 3, 'Sources: the LaTeX logs of tools/make_docs.py. "First build of 28 Oct": the first full build of the documents on 28 '
         'October 2026, before commit 0654dcd; "after 0654dcd": the rebuild of a fresh clone of 0654dcd; "now": the rebuild of a fresh clone of '
         '%s. Pages: pdfinfo. A loose line is an underfull line (TeX badness above 1000); an overflowing line is an overfull one.'
         % S['rebuild']['commit'], 11)
    wd.freeze_panes = 'B2'

    # ---------------------------------------------------------------- Fixes
    wf = wb.create_sheet('Fixes')
    header(wf, 1, ['#', 'Document', 'Section of the document', 'How', 'Change', 'Commit'], [5, 36, 40, 13, 60, 10])
    for i, f in enumerate(S['fixes'], 2):
        put(wf, i, [i - 1, f['document'], f['section'], f['how'], f['change'], f['commit']])
        wf.cell(row=i, column=1).alignment = CENTER; wf.cell(row=i, column=6).alignment = CENTER
    fl = len(S['fixes']) + 1
    note(wf, fl + 2, 'Every rewritten paragraph keeps its content: long inline formulas were moved into displayed equations, or the sentence '
         'was reworded so that TeX can break its lines evenly.', 6)
    wf.freeze_panes = 'A2'

    # ---------------------------------------------------------------- Checks
    wk = wb.create_sheet('Checks')
    header(wk, 1, ['Program', 'What it checks', 'Exit code', 'Results reported True', 'Results reported False', 'Verdict'], [46, 58, 10, 12, 12, 26])
    for i, p in enumerate(S['programs'], 2):
        put(wk, i, [p['program'], p['checks'], p['exit_code'], p['true'], p['false'], p['verdict']])
        for j in (3, 4, 5):
            wk.cell(row=i, column=j).alignment = CENTER
    pl = len(S['programs']) + 1
    put(wk, pl + 1, ['Total (%d programs)' % len(S['programs']),
                     '=COUNTIF(C2:C%d,0)&" of %d programs exit with code 0"' % (pl, len(S['programs'])), None,
                     '=SUM(D2:D%d)' % pl, '=SUM(E2:E%d)' % pl, None], bold=True)
    total_style(wk, pl + 1, 6)
    for j in (3, 4, 5):
        wk.cell(row=pl + 1, column=j).alignment = CENTER
    r0 = pl + 3
    block_title(wk, r0, 'The notebook 7+1.nb')
    nb = S['notebook']
    nb_rows = [('Checks in the notebook', nb['checks']), ('Checks that returned True', nb['true']), ('… printed where they stand', nb['printed']),
               ('… made by frameChecks (section 13.1), printed by the summary table', nb['frame']), ('Message (error) cells', nb['messages']),
               ('Input cells', nb['inputs']), ('Output cells', nb['outputs']), ('Evaluation time (seconds)', nb['seconds']),
               ('Wolfram version', nb['wolfram'])]
    for k, (a, v) in enumerate(nb_rows, r0 + 1):
        put(wk, k, [a, v], fmt={2: '0.000' if isinstance(v, float) else 'General'})
        wk.cell(row=k, column=2).alignment = LEFT
    nb_first = r0 + 1
    r1 = r0 + len(nb_rows) + 2
    fr = S['rebuild']
    block_title(wk, r1, 'Fresh clone of %s, rebuilt from scratch' % fr['commit'])
    fr_rows = [('Files tracked in the repository', fr['tracked']), ('Markdown and LaTeX files compared', fr['mdtex']),
               ('… identical byte for byte', fr['mdtex_same']), ('PDFs compared', fr['pdfs']),
               ('… with the same text and page count', fr['pdfs_same']), ('Documents with layout problems', fr['layout_problems']),
               ('Other files changed by the rebuild', fr['other_changed'])]
    for k, (a, v) in enumerate(fr_rows, r1 + 1):
        put(wk, k, [a, v]); wk.cell(row=k, column=2).alignment = LEFT
    fr_first = r1 + 1
    note(wk, r1 + len(fr_rows) + 2, 'Sources: the outputs of the programs in the rebuild of a fresh clone (build/docs/runs.json); "Results '
         'reported True" counts the lines of each output that end in ": True" or in "(True)", "Results reported False" those that end in ": False" '
         'or in "(False)". The notebook was evaluated on %s (tools/eval_log_7p1.txt).'
         % nb['finished'], 6)
    wk.freeze_panes = 'A2'

    # ---------------------------------------------------------------- Private files
    wp = wb.create_sheet('Private files')
    header(wp, 1, ['File', 'What it is', 'On GitHub?', 'How it is kept out'], [44, 46, 11, 50])
    for i, p in enumerate(S['private_files'], 2):
        put(wp, i, [p['file'], p['what'], p['on_github'], p['how']]); wp.cell(row=i, column=3).alignment = CENTER
    ql = len(S['private_files']) + 1
    note(wp, ql + 2, 'Source: git ls-files of a fresh clone of %s.' % fr['commit'], 4)
    wp.freeze_panes = 'A2'

    # ---------------------------------------------------------------- Summary (first sheet)
    ws = wb.create_sheet('Summary', 0)
    ws['A1'] = S['title']; ws['A1'].font = TITLE
    ws['A2'] = '%s, branch %s: the state at commit %s (%s)' % (S['repository'], S['branch'], dc, S['described_time']); ws['A2'].font = SUB
    header(ws, 4, ['Item', 'Value', 'Where it comes from'], [52, 16, 62])
    D = dl + 1
    rows = [
        ('Commit described (the latest when the spreadsheet was made)', '=Commits!A%d' % cl, 'Commits, last row'),
        ('Commits of 28 October 2026 up to the commit described', '=COUNTA(Commits!A2:A%d)' % cl, 'Commits'),
        ('File changes in these commits', '=Commits!E%d' % (cl + 1), 'Commits, total (a file changed twice counts twice)'),
        ('Documents (Markdown, LaTeX and PDF each)', '=COUNTA(Documents!A2:A%d)' % dl, 'Documents'),
        ('PDF pages', '=Documents!D%d' % D, 'Documents, total'),
        ('Loose lines in the first build of 28 October', '=Documents!E%d' % D, 'Documents, total'),
        ('Loose lines after commit 0654dcd', '=Documents!F%d' % D, 'Documents, total'),
        ('Loose lines now', '=Documents!G%d' % D, 'Documents, total'),
        ('Overflowing lines now', '=Documents!H%d' % D, 'Documents, total'),
        ('LaTeX errors now', '=Documents!I%d' % D, 'Documents, total'),
        ('Paragraphs rewritten (display, code block or rewording)',
         '=COUNTIF(Fixes!D2:D%d,"Display")+COUNTIF(Fixes!D2:D%d,"Code block")+COUNTIF(Fixes!D2:D%d,"Reworded")' % (fl, fl, fl), 'Fixes'),
        ('Check programs', '=COUNTA(Checks!A2:A%d)' % pl, 'Checks'),
        ('Check programs that exit with 0', '=COUNTIF(Checks!C2:C%d,0)' % pl, 'Checks'),
        ('Results reported True by the programs', '=Checks!D%d' % (pl + 1), 'Checks, total'),
        ('Results reported False by the programs', '=Checks!E%d' % (pl + 1), 'Checks, total'),
        ('Checks in the notebook 7+1.nb', '=Checks!B%d' % nb_first, 'Checks, the notebook'),
        ('Notebook checks that returned True', '=Checks!B%d' % (nb_first + 1), 'Checks, the notebook'),
        ('Notebook message (error) cells', '=Checks!B%d' % (nb_first + 4), 'Checks, the notebook'),
        ('Markdown and LaTeX files identical after a fresh rebuild', '=Checks!B%d&" of "&Checks!B%d' % (fr_first + 2, fr_first + 1),
         'Checks, fresh clone'),
        ('PDFs with the same text and page count after a fresh rebuild', '=Checks!B%d&" of "&Checks!B%d' % (fr_first + 4, fr_first + 3),
         'Checks, fresh clone'),
        ('Private files on GitHub', '=COUNTIF(\'Private files\'!C2:C%d,"Yes")' % ql, 'Private files'),
    ]
    for i, (a, f, src) in enumerate(rows, 5):
        put(ws, i, [a, f, src])
        ws.cell(row=i, column=2).alignment = Alignment(horizontal='right', vertical='top')
        ws.cell(row=i, column=3).font = GREY
    note(ws, len(rows) + 6, 'Every value in column B is a formula that reads the other sheets.', 3)
    ws.freeze_panes = 'A5'

    for sh in wb.worksheets:
        sh.page_setup.orientation = 'landscape'                     # printing: one page wide
        sh.page_setup.fitToWidth = 1
        sh.page_setup.fitToHeight = 0
        sh.sheet_properties.pageSetUpPr.fitToPage = True
    return wb


def recalculate(path):
    """let LibreOffice compute and store every formula value (headless open and save); False if soffice is not installed"""
    soffice = shutil.which('soffice') or shutil.which('libreoffice')
    if not soffice:
        return False
    with tempfile.TemporaryDirectory() as tmp:
        prof = 'file://' + os.path.join(tmp, 'profile')
        out = os.path.join(tmp, 'out')
        subprocess.run([soffice, '-env:UserInstallation=' + prof, '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML',
                        '--outdir', out, path], capture_output=True, text=True, timeout=180)
        res = os.path.join(out, os.path.basename(path))
        if not os.path.exists(res):
            raise SystemExit('LibreOffice did not write %s' % res)
        shutil.copyfile(res, path)
    return True


def main():
    S = json.load(open(DATA, encoding='utf-8'))
    build(S).save(OUT)
    done = recalculate(OUT)
    wb = load_workbook(OUT, data_only=True)
    formulas = load_workbook(OUT)
    n = errors = 0
    for ws in formulas.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    n += 1
                    v = wb[ws.title][c.coordinate].value
                    if done and (v is None or (isinstance(v, str) and v.startswith('#'))):
                        errors += 1; print('formula error:', ws.title, c.coordinate, c.value, '->', v)
    print('wrote %s: %d formulas, %s' % (OUT, n, ('%d errors after the LibreOffice recalculation' % errors) if done else
                                          'not recalculated (LibreOffice is not installed; Excel computes the values on opening)'))
    if done:
        for r in wb['Summary'].iter_rows(min_row=5, max_row=25, values_only=True):
            if r[0]:
                print('  %-62s %s' % (r[0], r[1]))
    return errors == 0


if __name__ == '__main__':
    sys.exit(0 if main() else 1)
